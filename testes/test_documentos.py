import io
import sqlite3

import pytest

from processamento import documentos, osc


def _pdf_com_texto(texto: str) -> bytes:
    """PDF mínimo válido de 1 página com texto (sem depender de biblioteca de escrita de PDF)."""
    conteudo = f"BT /F1 12 Tf 72 720 Td ({texto}) Tj ET".encode("latin-1")
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(conteudo)).encode() + b" >>\nstream\n" + conteudo + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    saida = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objetos, start=1):
        offsets.append(len(saida))
        saida += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    inicio_xref = len(saida)
    saida += f"xref\n0 {len(objetos) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        saida += f"{off:010d} 00000 n \n".encode()
    saida += f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\nstartxref\n{inicio_xref}\n%%EOF\n".encode()
    return bytes(saida)


def _docx_com_texto(paragrafos: list[str]) -> bytes:
    import docx

    d = docx.Document()
    for p in paragrafos:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx_com_linhas(linhas: list[list]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Projetos"
    for l in linhas:
        ws.append(l)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def conexao():
    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    osc.criar_tabelas(conn)
    yield conn
    conn.close()


@pytest.fixture
def osc_id(conexao):
    return osc.criar_ou_atualizar_osc(conexao, None, {"nome": "OSC Teste"})


def test_extrair_txt_utf8_e_latin1():
    assert documentos.extrair_texto("Educação pela arte".encode("utf-8"), ".txt")[0] == "Educação pela arte"
    assert documentos.extrair_texto("Educação".encode("latin-1"), ".txt")[0] == "Educação"


def test_extrair_pdf_real():
    texto, paginas = documentos.extrair_texto(_pdf_com_texto("Instituto Oswaldo Ribeiro de Mendonca"), ".pdf")
    assert "Oswaldo Ribeiro" in texto and paginas == 1


def test_extrair_docx_inclui_paragrafos_e_tabelas():
    import docx

    d = docx.Document()
    d.add_paragraph("Estatuto social do instituto")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Presidente"
    t.rows[0].cells[1].text = "Fulano de Tal"
    buf = io.BytesIO()
    d.save(buf)
    texto, _ = documentos.extrair_texto(buf.getvalue(), ".docx")
    assert "Estatuto social" in texto and "Presidente | Fulano de Tal" in texto


def test_extrair_xlsx_lista_abas_e_celulas():
    texto, abas = documentos.extrair_texto(_xlsx_com_linhas([["Projeto", "Valor"], ["Usina da Dança", 1000]]), ".xlsx")
    assert "[Aba: Projetos]" in texto and "Usina da Dança | 1000" in texto and abas == 1


def test_formato_desconhecido_levanta_erro():
    with pytest.raises(ValueError):
        documentos.extrair_texto(b"x", ".exe")


def test_adicionar_preserva_arquivo_original_e_extrai_texto(conexao, osc_id, tmp_path):
    conteudo = _docx_com_texto(["Relatório de impacto 2025", "Atendemos 800 crianças em Guaíra."])
    r = documentos.adicionar(conexao, osc_id, "Relatório Impacto 2025.docx", conteudo, tmp_path, "Relatório de impacto")
    assert r["status"] == documentos.STATUS_PROCESSADO and not r["duplicado"]
    doc = documentos.listar(conexao, osc_id)[0]
    assert doc["nome_arquivo"] == "Relatório Impacto 2025.docx"
    assert doc["extensao"] == ".docx" and doc["tamanho_bytes"] == len(conteudo)
    assert doc["categoria"] == "Relatório de impacto" and doc["enviado_em"]
    from pathlib import Path

    assert Path(doc["caminho_arquivo"]).read_bytes() == conteudo  # original intacto


def test_mesmo_arquivo_duas_vezes_nao_duplica(conexao, osc_id, tmp_path):
    conteudo = b"Texto institucional"
    documentos.adicionar(conexao, osc_id, "a.txt", conteudo, tmp_path)
    r2 = documentos.adicionar(conexao, osc_id, "copia_de_a.txt", conteudo, tmp_path)
    assert r2["duplicado"] is True
    assert len(documentos.listar(conexao, osc_id)) == 1


def test_pdf_sem_texto_fica_com_status_honesto(conexao, osc_id, tmp_path):
    r = documentos.adicionar(conexao, osc_id, "escaneado.pdf", _pdf_com_texto(""), tmp_path)
    assert r["status"] == documentos.STATUS_SEM_TEXTO
    assert documentos.obter_texto(conexao, r["id"]) is None


def test_arquivo_corrompido_registra_erro_sem_derrubar(conexao, osc_id, tmp_path):
    r = documentos.adicionar(conexao, osc_id, "quebrado.pdf", b"isto nao e um pdf", tmp_path)
    assert r["status"] == documentos.STATUS_ERRO
    assert documentos.listar(conexao, osc_id)[0]["detalhe_processamento"]


def test_formato_nao_suportado_e_guardado_mas_sem_leitura(conexao, osc_id, tmp_path):
    r = documentos.adicionar(conexao, osc_id, "foto.png", b"\x89PNG....", tmp_path)
    assert r["status"] == documentos.STATUS_NAO_SUPORTADO


def test_arquivo_vazio_e_rejeitado(conexao, osc_id, tmp_path):
    with pytest.raises(ValueError):
        documentos.adicionar(conexao, osc_id, "vazio.txt", b"", tmp_path)


def test_arquivo_grande_demais_e_rejeitado(conexao, osc_id, tmp_path, monkeypatch):
    monkeypatch.setattr(documentos, "TAMANHO_MAXIMO_BYTES", 10)
    with pytest.raises(ValueError):
        documentos.adicionar(conexao, osc_id, "grande.txt", b"x" * 11, tmp_path)


def test_busca_ignora_acento_e_devolve_trecho_literal_com_fonte(conexao, osc_id, tmp_path):
    documentos.adicionar(conexao, osc_id, "estatuto.txt",
                         "O instituto atua em Guaíra, Ipuã e Orlândia com educação através da arte.".encode(), tmp_path, "Estatuto")
    documentos.adicionar(conexao, osc_id, "outro.txt", b"Nada sobre o assunto aqui.", tmp_path)
    achados = documentos.buscar(conexao, osc_id, "EDUCACAO")
    assert len(achados) == 1
    assert achados[0]["documento"] == "estatuto.txt" and achados[0]["categoria"] == "Estatuto"
    assert "educação através da arte" in achados[0]["trechos"][0]  # trecho literal, com acento original


def test_busca_sem_resultado_e_termo_vazio(conexao, osc_id, tmp_path):
    documentos.adicionar(conexao, osc_id, "a.txt", b"conteudo qualquer", tmp_path)
    assert documentos.buscar(conexao, osc_id, "inexistente") == []
    assert documentos.buscar(conexao, osc_id, "   ") == []


def test_excluir_remove_registro_e_arquivo(conexao, osc_id, tmp_path):
    documentos.adicionar(conexao, osc_id, "a.txt", b"conteudo", tmp_path)
    doc = documentos.listar(conexao, osc_id)[0]
    from pathlib import Path

    assert Path(doc["caminho_arquivo"]).exists()
    assert documentos.excluir(conexao, doc["id"]) is True
    assert documentos.listar(conexao, osc_id) == []
    assert not Path(doc["caminho_arquivo"]).exists()
    assert documentos.excluir(conexao, doc["id"]) is False


def test_textos_para_contexto_so_traz_documentos_com_texto(conexao, osc_id, tmp_path):
    documentos.adicionar(conexao, osc_id, "com_texto.txt", b"Missao: educacao", tmp_path, "Estatuto")
    documentos.adicionar(conexao, osc_id, "escaneado.pdf", _pdf_com_texto(""), tmp_path)
    ctx = documentos.textos_para_contexto(conexao, osc_id)
    assert [c["documento"] for c in ctx] == ["com_texto.txt"]


def test_migracao_preserva_documento_antigo_sem_arquivo(conexao, osc_id):
    osc.adicionar_documento(conexao, osc_id, "estatuto", "Estatuto antigo", "descrição", "http://x")
    documentos.migrar(conexao)
    documentos.migrar(conexao)  # idempotente
    docs = documentos.listar(conexao, osc_id)
    assert docs[0]["nome"] == "Estatuto antigo" and docs[0]["caminho_arquivo"] is None
